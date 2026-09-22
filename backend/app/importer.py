"""Moteur d'import en masse des TI (tests d'intrusion / audits).

Le fichier fourni par l'analyste (CSV, TSV ou XLSX) est :
1. **lu et normalisé** (encodage, séparateur, en-têtes),
2. **analysé** : chaque colonne est rapprochée d'un champ de l'application ;
   les colonnes inconnues deviennent des *champs personnalisés* candidats et
   les valeurs référencées inconnues (catégorie, société, pilote, étiquette)
   sont remontées pour que l'analyste complète les détails manquants,
3. **importé** une fois les décisions saisies.

Aucune écriture en base n'est faite pendant l'analyse : tout est stocké dans le
lot d'import (`ImportBatch`) et l'import n'est effectif qu'à la validation.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from datetime import date, datetime
from difflib import SequenceMatcher

from app.models import AuditPriority, AuditStatus, CustomFieldType

MAX_PREVIEW_ROWS = 50
MAX_ROWS = 5000

# --- Champs cibles de l'application -----------------------------------------
# clé -> (libellé, type de résolution)
TARGET_FIELDS: dict[str, tuple[str, str]] = {
    "reference": ("Référence / identifiant TI", "scalar"),
    "name": ("Nom de l'audit", "scalar"),
    "description": ("Description", "scalar"),
    "category": ("Catégorie", "category"),
    "tags": ("Étiquettes (séparées par ;)", "tags"),
    "priority": ("Priorité", "enum_priority"),
    "status": ("Statut", "enum_status"),
    "pilot": ("Pilote d'audit", "user"),
    "service_owner": ("Responsable de service", "user"),
    "contact_name": ("Contact (nom)", "scalar"),
    "contact_email": ("Contact (email)", "scalar"),
    "contact_phone": ("Contact (téléphone)", "scalar"),
    "prestation_company": ("Société de prestation", "prestation_company"),
    "planned_start": ("Début prévu", "date"),
    "planned_end": ("Fin prévue", "date"),
    "actual_start": ("Début réel", "date"),
    "actual_end": ("Fin réelle", "date"),
    "estimated_days": ("Charge estimée (jours)", "number"),
}

ENTITY_TARGETS = {
    "category": "category",
    "prestation_company": "prestation_company",
    "pilot": "user",
    "service_owner": "user",
    "tags": "tag",
}

# Synonymes d'en-têtes reconnus automatiquement (après normalisation)
HEADER_SYNONYMS: dict[str, str] = {
    "reference": "reference",
    "ref": "reference",
    "id": "reference",
    "identifiant": "reference",
    "code": "reference",
    "code ti": "reference",
    "ti": "reference",
    "numero ti": "reference",
    "nom": "name",
    "nom audit": "name",
    "nom de l audit": "name",
    "libelle": "name",
    "intitule": "name",
    "titre": "name",
    "audit": "name",
    "test intrusion": "name",
    "nom du test": "name",
    "nom du ti": "name",
    "intitule du test": "name",
    "objet du test": "name",
    "nom de l audit": "name",
    "description": "description",
    "objet": "description",
    "perimetre": "description",
    "commentaire": "description",
    "categorie": "category",
    "type": "category",
    "type audit": "category",
    "type d audit": "category",
    "famille": "category",
    "tag": "tags",
    "tags": "tags",
    "etiquette": "tags",
    "etiquettes": "tags",
    "mots cles": "tags",
    "priorite": "priority",
    "criticite": "priority",
    "statut": "status",
    "etat": "status",
    "avancement": "status",
    "pilote": "pilot",
    "pilote audit": "pilot",
    "pilote d audit": "pilot",
    "responsable audit": "pilot",
    "charge de suivi": "pilot",
    "responsable service": "service_owner",
    "responsable de service": "service_owner",
    "proprietaire": "service_owner",
    "moa": "service_owner",
    "contact": "contact_name",
    "contact nom": "contact_name",
    "nom contact": "contact_name",
    "email": "contact_email",
    "mail": "contact_email",
    "email contact": "contact_email",
    "adresse mail": "contact_email",
    "telephone": "contact_phone",
    "tel": "contact_phone",
    "prestataire": "prestation_company",
    "societe": "prestation_company",
    "societe de prestation": "prestation_company",
    "fournisseur": "prestation_company",
    "cabinet": "prestation_company",
    "date debut": "planned_start",
    "debut": "planned_start",
    "debut prevu": "planned_start",
    "date debut prevue": "planned_start",
    "date de debut": "planned_start",
    "date fin": "planned_end",
    "fin": "planned_end",
    "fin prevue": "planned_end",
    "date fin prevue": "planned_end",
    "date de fin": "planned_end",
    "debut reel": "actual_start",
    "date debut reelle": "actual_start",
    "date reelle debut": "actual_start",
    "fin reelle": "actual_end",
    "date fin reelle": "actual_end",
    "date reelle fin": "actual_end",
    "charge": "estimated_days",
    "charge estimee": "estimated_days",
    "jours": "estimated_days",
    "nb jours": "estimated_days",
    "jours homme": "estimated_days",
}

PRIORITY_SYNONYMS = {
    "basse": AuditPriority.BASSE, "faible": AuditPriority.BASSE, "low": AuditPriority.BASSE, "p4": AuditPriority.BASSE,
    "moyenne": AuditPriority.MOYENNE, "moyen": AuditPriority.MOYENNE, "normale": AuditPriority.MOYENNE,
    "medium": AuditPriority.MOYENNE, "p3": AuditPriority.MOYENNE,
    "haute": AuditPriority.HAUTE, "elevee": AuditPriority.HAUTE, "high": AuditPriority.HAUTE, "p2": AuditPriority.HAUTE,
    "critique": AuditPriority.CRITIQUE, "critical": AuditPriority.CRITIQUE, "majeure": AuditPriority.CRITIQUE,
    "p1": AuditPriority.CRITIQUE,
}

STATUS_SYNONYMS = {
    "brouillon": AuditStatus.BROUILLON, "draft": AuditStatus.BROUILLON, "a qualifier": AuditStatus.BROUILLON,
    "planifie": AuditStatus.PLANIFIE, "planifiee": AuditStatus.PLANIFIE, "a planifier": AuditStatus.PLANIFIE,
    "prevu": AuditStatus.PLANIFIE, "planned": AuditStatus.PLANIFIE,
    "en cours": AuditStatus.EN_COURS, "demarre": AuditStatus.EN_COURS, "in progress": AuditStatus.EN_COURS,
    "en attente": AuditStatus.EN_ATTENTE, "suspendu": AuditStatus.EN_ATTENTE, "on hold": AuditStatus.EN_ATTENTE,
    "bloque": AuditStatus.BLOQUE, "blocked": AuditStatus.BLOQUE,
    "termine": AuditStatus.TERMINE, "terminee": AuditStatus.TERMINE, "clos": AuditStatus.TERMINE,
    "cloture": AuditStatus.TERMINE, "done": AuditStatus.TERMINE, "fini": AuditStatus.TERMINE,
    "annule": AuditStatus.ANNULE, "abandonne": AuditStatus.ANNULE, "cancelled": AuditStatus.ANNULE,
}

TRUE_VALUES = {"oui", "yes", "true", "vrai", "1", "x"}
FALSE_VALUES = {"non", "no", "false", "faux", "0"}

DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y/%m/%d", "%d/%m/%y", "%m/%d/%Y", "%Y%m%d")


# ---------------------------------------------------------------------------
# Normalisation
# ---------------------------------------------------------------------------
def normalize(value: str | None) -> str:
    """Minuscule, sans accent ni ponctuation : sert aux rapprochements."""
    if value is None:
        return ""
    text = unicodedata.normalize("NFKD", str(value))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^a-zA-Z0-9]+", " ", text).strip().lower()
    return re.sub(r"\s+", " ", text)


def slugify_key(value: str) -> str:
    key = re.sub(r"[^a-z0-9]+", "_", normalize(value)).strip("_")
    return (key or "champ")[:80]


def parse_date(value) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    text = text.split(" ")[0] if " " in text and ":" in text else text
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


# Un nombre commence par un chiffre (ou un signe) : « APP0 » n'est pas un nombre,
# « 12 j » et « 1 234,5 » en sont.
NUMBER_RE = re.compile(r"^[+-]?\d[\d\s]*(?:[.,]\d+)?")


def parse_number(value) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    match = NUMBER_RE.match(str(value).strip())
    if match is None:
        return None
    cleaned = re.sub(r"\s", "", match.group(0)).replace(",", ".")
    try:
        return float(cleaned)
    except ValueError:
        return None


def parse_bool(value) -> bool | None:
    text = normalize(value)
    if text in TRUE_VALUES:
        return True
    if text in FALSE_VALUES:
        return False
    return None


def split_multi(value: str | None) -> list[str]:
    if not value:
        return []
    parts = re.split(r"[;,/|]", str(value))
    return [p.strip() for p in parts if p.strip()]


# ---------------------------------------------------------------------------
# Lecture de fichier
# ---------------------------------------------------------------------------
def read_table(filename: str, content: bytes) -> tuple[list[str], list[dict[str, str | None]]]:
    """Retourne (colonnes, lignes) à partir d'un CSV/TSV/XLSX."""
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xlsm")):
        return _read_xlsx(content)
    return _read_csv(content)


def _read_csv(content: bytes) -> tuple[list[str], list[dict[str, str | None]]]:
    text: str | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("Encodage du fichier non reconnu (UTF-8 ou Windows-1252 attendus)")

    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","

    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    raw_rows = [row for row in reader if any((cell or "").strip() for cell in row)]
    if not raw_rows:
        raise ValueError("Le fichier ne contient aucune ligne exploitable")

    headers = _dedupe_headers([(cell or "").strip() for cell in raw_rows[0]])
    rows: list[dict[str, str | None]] = []
    for raw in raw_rows[1 : MAX_ROWS + 1]:
        rows.append({headers[i]: (raw[i].strip() if i < len(raw) and raw[i] is not None else None) for i in range(len(headers))})
    return headers, rows


def _read_xlsx(content: bytes) -> tuple[list[str], list[dict[str, str | None]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ValueError("Le support XLSX nécessite la bibliothèque openpyxl") from exc

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    headers: list[str] | None = None
    rows: list[dict[str, str | None]] = []
    for raw in iterator:
        if raw is None or all(cell is None or str(cell).strip() == "" for cell in raw):
            continue
        if headers is None:
            headers = _dedupe_headers([str(cell).strip() if cell is not None else "" for cell in raw])
            continue
        values: dict[str, str | None] = {}
        for i, header in enumerate(headers):
            cell = raw[i] if i < len(raw) else None
            if isinstance(cell, datetime):
                values[header] = cell.date().isoformat()
            elif isinstance(cell, date):
                values[header] = cell.isoformat()
            elif cell is None:
                values[header] = None
            else:
                values[header] = str(cell).strip()
        rows.append(values)
        if len(rows) >= MAX_ROWS:
            break
    workbook.close()
    if headers is None:
        raise ValueError("Le fichier ne contient aucune ligne exploitable")
    return headers, rows


def _dedupe_headers(headers: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    result: list[str] = []
    for index, header in enumerate(headers):
        name = header or f"Colonne {index + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        result.append(name)
    return result


# ---------------------------------------------------------------------------
# Analyse
# ---------------------------------------------------------------------------
# Mots trop génériques pour trancher seuls entre deux champs cibles
GENERIC_TOKENS = {"nom", "date", "type", "audit", "code", "id", "contact", "etat", "numero"}


def _contained_synonym(key: str) -> str | None:
    """Cherche le synonyme le plus spécifique contenu dans l'en-tête.

    Permet de reconnaître « Nom du test », « Date de début prévue » ou
    « Nom du prestataire ». Arbitrage, dans l'ordre : le synonyme le plus long,
    puis un mot porteur de sens plutôt qu'un mot générique (« prestataire »
    l'emporte sur « nom »), puis le premier rencontré (« Type d'audit » reste
    une catégorie et ne devient pas un nom d'audit).
    """
    tokens = key.split()
    best: tuple[int, int, int, str] | None = None  # (longueur, spécificité, -position, cible)
    for synonym, candidate in HEADER_SYNONYMS.items():
        parts = synonym.split()
        length = len(parts)
        if length > len(tokens):
            continue
        for start in range(len(tokens) - length + 1):
            if tokens[start : start + length] != parts:
                continue
            specific = 0 if (length == 1 and parts[0] in GENERIC_TOKENS) else 1
            score = (length, specific, -start, candidate)
            if best is None or score[:3] > best[:3]:
                best = score
            break
    return best[3] if best else None


def suggest_mapping(headers: list[str]) -> dict[str, str]:
    """Rapproche chaque en-tête d'un champ cible.

    Trois passes successives : correspondance exacte d'un synonyme, forte
    similarité textuelle (« Fin réel » ≈ « fin reelle »), puis synonyme contenu
    dans l'intitulé (« Nom du test » → nom de l'audit). Un champ cible n'est
    proposé qu'une fois (sauf les étiquettes, qui peuvent venir de plusieurs
    colonnes). Le résultat reste modifiable par l'analyste.
    """
    mapping: dict[str, str] = {}
    used: set[str] = set()
    for header in headers:
        key = normalize(header)
        target = HEADER_SYNONYMS.get(key)
        if target is None:
            best, score = None, 0.0
            for synonym, candidate in HEADER_SYNONYMS.items():
                ratio = SequenceMatcher(None, key, synonym).ratio()
                if ratio > score:
                    best, score = candidate, ratio
            target = best if score >= 0.86 else _contained_synonym(key)
        if target and target not in used:
            mapping[header] = target
            if target != "tags":
                used.add(target)
    return mapping


def detect_type(values: list[str | None]) -> CustomFieldType:
    filled = [v for v in values if v not in (None, "")]
    if not filled:
        return CustomFieldType.TEXTE
    if all(parse_date(v) is not None for v in filled):
        return CustomFieldType.DATE
    if all(parse_number(v) is not None for v in filled):
        return CustomFieldType.NOMBRE
    if all(normalize(v) in TRUE_VALUES | FALSE_VALUES for v in filled):
        return CustomFieldType.BOOLEEN
    distinct = {str(v).strip() for v in filled}
    if len(distinct) <= max(8, len(filled) // 10) and len(distinct) < len(filled):
        return CustomFieldType.LISTE
    if max(len(str(v)) for v in filled) > 120:
        return CustomFieldType.TEXTE_LONG
    return CustomFieldType.TEXTE


def best_match(value: str, candidates: dict[str, str], threshold: float = 0.84) -> tuple[str | None, str | None]:
    """Cherche l'enregistrement existant le plus proche (id, libellé)."""
    target = normalize(value)
    best_id, best_label, best_score = None, None, 0.0
    for identifier, label in candidates.items():
        score = SequenceMatcher(None, target, normalize(label)).ratio()
        if normalize(label) == target:
            return identifier, label
        if score > best_score:
            best_id, best_label, best_score = identifier, label, score
    if best_score >= threshold:
        return best_id, best_label
    return None, None


def coerce_value(target: str, raw):
    """Convertit une valeur brute vers le type attendu par le champ cible."""
    kind = TARGET_FIELDS.get(target, ("", "scalar"))[1]
    if raw in (None, ""):
        return None
    if kind == "date":
        return parse_date(raw)
    if kind == "number":
        return parse_number(raw)
    if kind == "enum_priority":
        return PRIORITY_SYNONYMS.get(normalize(raw))
    if kind == "enum_status":
        return STATUS_SYNONYMS.get(normalize(raw))
    if kind == "tags":
        return split_multi(raw)
    return str(raw).strip()
